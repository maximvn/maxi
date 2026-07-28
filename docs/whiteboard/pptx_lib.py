"""Bouwstenen voor de PowerPoint: alles wordt als echte vorm en echt tekstkader
geplaatst, zodat je in PowerPoint elke tekst kunt selecteren en aanpassen.

De tekst wordt vooraf opgemeten (met DejaVu Sans als conservatieve maatstaf voor
Calibri, dat smaller is) en zo nodig een maatje kleiner gezet, zodat er niets
buiten zijn kader valt.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from svg_lib import afbreken, meet

# ------------------------------------------------------------------ palet
INKT = "0F1B2D"
NAVY = "123A63"
BLAUW = "1F6FB2"
TEAL = "0E9A87"
AMBER = "DE8A1B"
ROOD = "C4392D"
PAARS = "6B4BA8"
LEI = "64748B"
LIJN = "D8E0EA"
ACHTER = "F4F7FB"
WIT = "FFFFFF"

LETTER = "Calibri"
LETTER_KOP = "Calibri"

BREEDTE_DIA = 13.333
HOOGTE_DIA = 7.5
MARGE = 0.55
INHOUD_B = BREEDTE_DIA - 2 * MARGE
INHOUD_Y = 1.58
INHOUD_H = 6.95 - INHOUD_Y


def rgb(hexkleur):
    return RGBColor.from_string(hexkleur)


def licht(hexkleur, f=0.12):
    r, g, b = (int(hexkleur[i:i + 2], 16) for i in (0, 2, 4))
    r = round(255 + (r - 255) * f)
    g = round(255 + (g - 255) * f)
    b = round(255 + (b - 255) * f)
    return f"{r:02X}{g:02X}{b:02X}"


# --------------------------------------------------------------- meten
def _regels(tekst, size, breedte_pt, vet):
    uit = []
    for stuk in tekst.split("\n"):
        uit.extend(afbreken(stuk, size, breedte_pt, "bold" if vet else "normal")
                   or [""])
    return uit


def passende_maat(tekst, breedte_in, hoogte_in, start=16, minimum=8, vet=False,
                  lh=1.22, marge_pt=8):
    """Grootste puntgrootte waarbij de tekst binnen breedte en hoogte blijft."""
    bp = breedte_in * 72 - marge_pt
    hp = hoogte_in * 72 - 4
    size = start
    while size > minimum:
        regels = _regels(tekst, size, bp, vet)
        if len(regels) * size * lh <= hp:
            return size
        size -= 0.5
    return minimum


def regelhoogte(tekst, breedte_in, size, vet=False, lh=1.22, marge_pt=8):
    bp = breedte_in * 72 - marge_pt
    return len(_regels(tekst, size, bp, vet)) * size * lh / 72.0


# ------------------------------------------------------------- basis
def nieuwe_presentatie():
    prs = Presentation()
    prs.slide_width = Inches(BREEDTE_DIA)
    prs.slide_height = Inches(HOOGTE_DIA)
    return prs


def lege_dia(prs, achtergrond=ACHTER):
    dia = prs.slides.add_slide(prs.slide_layouts[6])
    vlak(dia, 0, 0, BREEDTE_DIA, HOOGTE_DIA, achtergrond)
    return dia


def vlak(dia, x, y, w, h, vul, rand=None, randbreedte=1.0, radius=None):
    vorm = dia.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        try:
            vorm.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    if vul is None:
        vorm.fill.background()
    else:
        vorm.fill.solid()
        vorm.fill.fore_color.rgb = rgb(vul)
    if rand:
        vorm.line.color.rgb = rgb(rand)
        vorm.line.width = Pt(randbreedte)
    else:
        vorm.line.fill.background()
    vorm.shadow.inherit = False
    if vorm.has_text_frame:
        vorm.text_frame.text = ""
    return vorm


def kaart(dia, x, y, w, h, vul=WIT, rand=LIJN, radius=0.055, randbreedte=1.0):
    return vlak(dia, x, y, w, h, vul, rand, randbreedte, radius)


def cirkel(dia, cx, cy, d, vul, label=None, tekstkleur=WIT, size=14):
    vorm = dia.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2),
                                Inches(cy - d / 2), Inches(d), Inches(d))
    vorm.fill.solid()
    vorm.fill.fore_color.rgb = rgb(vul)
    vorm.line.fill.background()
    vorm.shadow.inherit = False
    tf = vorm.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label or ""
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = LETTER
    r.font.color.rgb = rgb(tekstkleur)
    return vorm


def stip(dia, cx, cy, kleur, d=0.13):
    ring = dia.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d * 0.85),
                                Inches(cy - d * 0.85), Inches(d * 1.7),
                                Inches(d * 1.7))
    ring.fill.solid()
    ring.fill.fore_color.rgb = rgb(licht(kleur, 0.22))
    ring.line.fill.background()
    ring.shadow.inherit = False
    kern = dia.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2),
                                Inches(cy - d / 2), Inches(d), Inches(d))
    kern.fill.solid()
    kern.fill.fore_color.rgb = rgb(kleur)
    kern.line.fill.background()
    kern.shadow.inherit = False
    return kern


def pijl(dia, x, y, w, h, kleur, richting="omlaag"):
    vorm = dia.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW if richting == "omlaag" else MSO_SHAPE.RIGHT_ARROW,
        Inches(x), Inches(y), Inches(w), Inches(h))
    vorm.fill.solid()
    vorm.fill.fore_color.rgb = rgb(kleur)
    vorm.line.fill.background()
    vorm.shadow.inherit = False
    return vorm


def pil(dia, x, y, w, h, label, vul, rand, tekstkleur=None, size=12, vet=True):
    vorm = dia.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    try:
        vorm.adjustments[0] = 0.5
    except (IndexError, ValueError):
        pass
    vorm.fill.solid()
    vorm.fill.fore_color.rgb = rgb(vul)
    if rand:
        vorm.line.color.rgb = rgb(rand)
        vorm.line.width = Pt(1.0)
    else:
        vorm.line.fill.background()
    vorm.shadow.inherit = False
    tf = vorm.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(passende_maat(label, w - 0.12, h, size, 8, vet))
    r.font.bold = vet
    r.font.name = LETTER
    r.font.color.rgb = rgb(tekstkleur or rand or INKT)
    return vorm


# ---------------------------------------------------------------- tekst
def tekst(dia, x, y, w, h, blokken, anchor="top", uitlijning="left",
          marge=0.06, autofit=True):
    """blokken: lijst van dicts met tekst, size, kleur, vet, cursief, na, lh."""
    if isinstance(blokken, str):
        blokken = [{"tekst": blokken}]
    doos = dia.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = doos.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(marge)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "midden": MSO_ANCHOR.MIDDLE,
                          "onder": MSO_ANCHOR.BOTTOM}[anchor]

    if autofit:
        beschikbaar = h - 0.06
        while True:
            totaal = 0.0
            for b in blokken:
                s = b.get("size", 14)
                totaal += regelhoogte(b["tekst"], w - 2 * marge, s,
                                      b.get("vet", False), b.get("lh", 1.2))
                totaal += b.get("na", 4) / 72.0
            if totaal <= beschikbaar or all(b.get("size", 14) <= 8 for b in blokken):
                break
            for b in blokken:
                b["size"] = max(8, b.get("size", 14) - 0.5)

    for i, b in enumerate(blokken):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                       "right": PP_ALIGN.RIGHT}[b.get("uitlijning", uitlijning)]
        p.space_after = Pt(b.get("na", 4))
        p.line_spacing = b.get("lh", 1.2)
        for j, regel in enumerate(b["tekst"].split("\n")):
            if j:
                p = tf.add_paragraph()
                p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                               "right": PP_ALIGN.RIGHT}[b.get("uitlijning", uitlijning)]
                p.space_after = Pt(b.get("na", 4))
                p.line_spacing = b.get("lh", 1.2)
            r = p.add_run()
            r.text = regel
            r.font.size = Pt(b.get("size", 14))
            r.font.bold = b.get("vet", False)
            r.font.italic = b.get("cursief", False)
            r.font.name = b.get("letter", LETTER)
            r.font.color.rgb = rgb(b.get("kleur", INKT))
    return doos


def kop(dia, nummer, titel, sub=None, kleur=NAVY):
    """Diakop met genummerde penning; het herhalende motief van de reeks."""
    if nummer is not None:
        cirkel(dia, MARGE + 0.25, 0.72, 0.5, kleur, str(nummer), WIT, 16)
        tx = MARGE + 0.66
    else:
        tx = MARGE
    tekst(dia, tx, 0.40, INHOUD_B - (tx - MARGE), 0.44,
          [{"tekst": titel, "size": 27, "vet": True, "kleur": INKT, "na": 0,
            "letter": LETTER_KOP}], anchor="midden", marge=0)
    if sub:
        tekst(dia, tx, 0.90, INHOUD_B - (tx - MARGE), 0.36,
              [{"tekst": sub, "size": 13.5, "kleur": LEI, "na": 0}],
              anchor="midden", marge=0)
    return INHOUD_Y


def voet(dia, nummer, label="Acute poort en Hotfloor"):
    tekst(dia, MARGE, 7.02, 6.0, 0.3,
          [{"tekst": label, "size": 9, "kleur": LEI, "na": 0}], marge=0)
    tekst(dia, BREEDTE_DIA - MARGE - 1.2, 7.02, 1.2, 0.3,
          [{"tekst": str(nummer), "size": 9, "kleur": LEI, "na": 0,
            "uitlijning": "right"}], marge=0)


def melding(dia, y, label, boodschap, kleur=AMBER, h=0.86):
    """Accentvlak met de vraag of waarschuwing die bij de dia hoort."""
    kaart(dia, MARGE, y, INHOUD_B, h, licht(kleur, 0.13), licht(kleur, 0.45))
    cirkel(dia, MARGE + 0.36, y + h / 2, 0.34, kleur, "!", WIT, 14)
    tekst(dia, MARGE + 0.66, y + 0.08, INHOUD_B - 0.9, h - 0.16,
          [{"tekst": label.upper(), "size": 9.5, "vet": True, "kleur": kleur, "na": 2},
           {"tekst": boodschap, "size": 13, "vet": True, "kleur": INKT, "na": 0}],
          anchor="midden")


# ---------------------------------------------------------------- tabel
def tabel(dia, x, y, w, kolombreedtes, koppen, rijen, kopkleur=NAVY,
          size=11, kopsize=10.5, rijhoogte=0.42, kophoogte=0.4,
          eerste_vet=True, kleuren=None, uitlijningen=None, rijhoogtes=None):
    """Echte PowerPoint-tabel: rijen, kolommen en tekst blijven bewerkbaar.

    Een waarde met regeleindes wordt in losse alinea's gezet, zodat een
    opsomming in een cel ook echt als opsomming leest.
    """
    n_rij = len(rijen) + (1 if koppen else 0)
    hoogtes = rijhoogtes or [rijhoogte] * len(rijen)
    vorm = dia.shapes.add_table(n_rij, len(kolombreedtes), Inches(x), Inches(y),
                                Inches(w), Inches(kophoogte + sum(hoogtes)))
    tbl = vorm.table
    tbl.first_row = bool(koppen)
    tbl.horz_banding = False
    for i, b in enumerate(kolombreedtes):
        tbl.columns[i].width = Inches(b)
    if uitlijningen is None:
        uitlijningen = ["left"] + ["center"] * (len(kolombreedtes) - 1)

    def vul_cel(cel, waarde, vet, kleur, size_, uitlijning, achtergrondkleur):
        cel.fill.solid()
        cel.fill.fore_color.rgb = rgb(achtergrondkleur)
        cel.margin_left = cel.margin_right = Inches(0.09)
        cel.margin_top = cel.margin_bottom = Inches(0.04)
        cel.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cel.text_frame
        tf.word_wrap = True
        for k, regel in enumerate(str(waarde).split("\n")):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                           "right": PP_ALIGN.RIGHT}[uitlijning]
            p.space_after = Pt(2)
            r = p.add_run()
            r.text = regel
            r.font.size = Pt(size_)
            r.font.bold = vet
            r.font.name = LETTER
            r.font.color.rgb = rgb(kleur)

    if koppen:
        tbl.rows[0].height = Inches(kophoogte)
        for j, k in enumerate(koppen):
            maat = passende_maat(str(k), kolombreedtes[j] - 0.2, kophoogte,
                                 kopsize, 7.5, True)
            vul_cel(tbl.cell(0, j), k, True, WIT, maat, uitlijningen[j], kopkleur)
    voor = 1 if koppen else 0
    for i, rij in enumerate(rijen):
        tbl.rows[i + voor].height = Inches(hoogtes[i])
        for j, waarde in enumerate(rij):
            kleur = INKT
            if kleuren and kleuren[i][j]:
                kleur = kleuren[i][j]
            regels = str(waarde).split("\n")
            maat = min(passende_maat(r, kolombreedtes[j] - 0.2,
                                     hoogtes[i] / max(1, len(regels)), size, 7.5,
                                     j == 0 and eerste_vet) for r in regels)
            vul_cel(tbl.cell(i + voor, j), waarde,
                    (j == 0 and eerste_vet) or
                    (kleuren and kleuren[i][j] in (ROOD, AMBER)),
                    kleur, maat, uitlijningen[j],
                    WIT if i % 2 == 0 else ACHTER)
    return vorm
