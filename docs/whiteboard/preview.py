#!/usr/bin/env python3
"""Rendert de PowerPoint naar afbeeldingen om de opmaak te kunnen controleren.

LibreOffice draait niet in deze omgeving, dus de dia's worden hier uit de
pptx-geometrie opnieuw getekend: posities, vullingen, randen en tekst. Het is
geen exacte PowerPoint-weergave, maar wel nauwkeurig genoeg om overloop,
overlap en uitlijning te zien.
"""
import os
import sys

import cairosvg
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu

from svg_lib import Svg, afbreken, meet

PPI = 100.0
PT = PPI / 72.0
HIER = os.path.dirname(os.path.abspath(__file__))
UIT = os.path.join(HIER, "preview")


def inch(waarde):
    return (waarde or 0) / 914400.0


def kleur_van(obj, standaard=None):
    try:
        if obj.type is None:
            return standaard
        return "#" + str(obj.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        return standaard


def lijnkleur(vorm):
    try:
        if vorm.line.fill.type is None or vorm.line.fill.type == 5:
            return None
        return "#" + str(vorm.line.color.rgb)
    except (AttributeError, TypeError, ValueError):
        return None


def teken_tekst(d, tf, x, y, w, h, waarschuw, naam):
    ml = inch(tf.margin_left) if tf.margin_left is not None else 0.1
    mr = inch(tf.margin_right) if tf.margin_right is not None else 0.1
    mt = inch(tf.margin_top) if tf.margin_top is not None else 0.05
    breedte = (w - ml - mr) * PPI
    if breedte <= 4:
        return

    blokken = []
    for p in tf.paragraphs:
        stukken = [(r.text, r) for r in p.runs if r.text]
        if not stukken:
            blokken.append((None, "", 12 * PT, None, False, 0, 1.2))
            continue
        tekst = "".join(s for s, _ in stukken)
        r0 = stukken[0][1]
        size = (r0.font.size.pt if r0.font.size else 18) * PT
        kleur = "#" + str(r0.font.color.rgb) if (r0.font.color and
                                                 r0.font.color.type is not None) else "#0F1B2D"
        vet = bool(r0.font.bold)
        na = (p.space_after.pt if p.space_after else 0) * PT
        lh = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        blokken.append((p.alignment, tekst, size, kleur, vet, na, lh))

    hoogte = 0.0
    uitgevouwen = []
    for uitl, tekst, size, kleur, vet, na, lh in blokken:
        regels = afbreken(tekst, size, breedte, "bold" if vet else "normal") or [""]
        uitgevouwen.append((uitl, regels, size, kleur, vet, na, lh))
        hoogte += len(regels) * size * lh + na

    beschikbaar = (h - mt * 2) * PPI
    anchor = tf.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE:
        cy = (y + h / 2) * PPI - hoogte / 2
    elif anchor == MSO_ANCHOR.BOTTOM:
        cy = (y + h) * PPI - mt * PPI - hoogte
    else:
        cy = (y + mt) * PPI

    if hoogte > beschikbaar + 1.5:
        waarschuw.append(f"{naam}: tekst {hoogte / PPI:.2f}in in vak van "
                         f"{beschikbaar / PPI:.2f}in -> {blokken[0][1][:48]!r}")

    for uitl, regels, size, kleur, vet, na, lh in uitgevouwen:
        for regel in regels:
            cy += size * lh
            basis = cy - size * 0.26
            if uitl == PP_ALIGN.CENTER:
                d.text((x + w / 2) * PPI, basis, regel, size, kleur, "middle",
                       "bold" if vet else "normal")
            elif uitl == PP_ALIGN.RIGHT:
                d.text((x + w - mr) * PPI, basis, regel, size, kleur, "end",
                       "bold" if vet else "normal")
            else:
                d.text((x + ml) * PPI, basis, regel, size, kleur, "start",
                       "bold" if vet else "normal")
        cy += na


def teken_vorm(d, vorm, waarschuw, prefix=""):
    if vorm.shape_type == MSO_SHAPE_TYPE.TABLE:
        teken_tabel(d, vorm, waarschuw, prefix)
        return
    x, y = inch(vorm.left), inch(vorm.top)
    w, h = inch(vorm.width), inch(vorm.height)
    naam = f"{prefix}{vorm.shape_id}"

    if vorm.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        vul = kleur_van(vorm.fill)
        rand = lijnkleur(vorm)
        soort = str(vorm.auto_shape_type)
        if "OVAL" in soort:
            d.parts.append(
                f'<ellipse cx="{(x + w / 2) * PPI:.1f}" cy="{(y + h / 2) * PPI:.1f}" '
                f'rx="{w * PPI / 2:.1f}" ry="{h * PPI / 2:.1f}" '
                f'fill="{vul or "none"}" stroke="{rand or "none"}" stroke-width="1"/>')
        elif "DOWN_ARROW" in soort:
            x0, y0, x1, y1 = x * PPI, y * PPI, (x + w) * PPI, (y + h) * PPI
            mx = (x0 + x1) / 2
            d.parts.append(
                f'<polygon points="{mx - w * PPI * 0.22:.1f},{y0:.1f} '
                f'{mx + w * PPI * 0.22:.1f},{y0:.1f} '
                f'{mx + w * PPI * 0.22:.1f},{y1 - h * PPI * 0.45:.1f} '
                f'{x1:.1f},{y1 - h * PPI * 0.45:.1f} {mx:.1f},{y1:.1f} '
                f'{x0:.1f},{y1 - h * PPI * 0.45:.1f} '
                f'{mx - w * PPI * 0.22:.1f},{y1 - h * PPI * 0.45:.1f}" '
                f'fill="{vul or "none"}"/>')
        elif "RIGHT_ARROW" in soort:
            x0, y0, x1, y1 = x * PPI, y * PPI, (x + w) * PPI, (y + h) * PPI
            my = (y0 + y1) / 2
            d.parts.append(
                f'<polygon points="{x0:.1f},{my - h * PPI * 0.22:.1f} '
                f'{x1 - w * PPI * 0.45:.1f},{my - h * PPI * 0.22:.1f} '
                f'{x1 - w * PPI * 0.45:.1f},{y0:.1f} {x1:.1f},{my:.1f} '
                f'{x1 - w * PPI * 0.45:.1f},{y1:.1f} '
                f'{x1 - w * PPI * 0.45:.1f},{my + h * PPI * 0.22:.1f} '
                f'{x0:.1f},{my + h * PPI * 0.22:.1f}" fill="{vul or "none"}"/>')
        else:
            straal = 10 if "ROUNDED" in soort else 0
            if "ROUNDED" in soort:
                try:
                    straal = min(w, h) * PPI * float(vorm.adjustments[0])
                except (IndexError, ValueError, TypeError):
                    straal = 10
            d.rect(x * PPI, y * PPI, w * PPI, h * PPI, fill=vul or "none",
                   stroke=rand or "none", sw=1, rx=straal)
    if vorm.has_text_frame and vorm.text_frame.text.strip():
        teken_tekst(d, vorm.text_frame, x, y, w, h, waarschuw, naam)


def teken_tabel(d, vorm, waarschuw, prefix=""):
    tbl = vorm.table
    x0, y0 = inch(vorm.left), inch(vorm.top)
    ky = y0
    for i, rij in enumerate(tbl.rows):
        kx = x0
        hoogte = inch(rij.height)
        for j, kolom in enumerate(tbl.columns):
            breedte = inch(kolom.width)
            cel = tbl.cell(i, j)
            vul = kleur_van(cel.fill, "#FFFFFF")
            d.rect(kx * PPI, ky * PPI, breedte * PPI, hoogte * PPI, fill=vul,
                   stroke="#C9D3DF", sw=1, rx=0)
            if cel.text.strip():
                teken_tekst(d, cel.text_frame, kx, ky, breedte, hoogte,
                            waarschuw, f"{prefix}tabel r{i}k{j}")
            kx += breedte
        ky += hoogte


def main():
    pad = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        HIER, "Acute-poort-en-Hotfloor.pptx")
    os.makedirs(UIT, exist_ok=True)
    for oud in os.listdir(UIT):
        os.remove(os.path.join(UIT, oud))

    prs = Presentation(pad)
    bw, bh = inch(prs.slide_width) * PPI, inch(prs.slide_height) * PPI
    alle = []
    for n, dia in enumerate(prs.slides, start=1):
        d = Svg(int(bw), int(bh), bg="#FFFFFF")
        waarschuw = []
        for vorm in dia.shapes:
            teken_vorm(d, vorm, waarschuw, f"dia{n} ")
            if (inch(vorm.left) < -0.02 or inch(vorm.top) < -0.02
                    or inch(vorm.left) + inch(vorm.width) > inch(prs.slide_width) + 0.02
                    or inch(vorm.top) + inch(vorm.height) > inch(prs.slide_height) + 0.02):
                waarschuw.append(f"dia {n}: vorm {vorm.shape_id} valt buiten de dia")
        svg = os.path.join(UIT, f"dia-{n:02d}.svg")
        d.save(svg)
        cairosvg.svg2png(url=svg, write_to=os.path.join(UIT, f"dia-{n:02d}.png"),
                         output_width=int(bw * 1.6), output_height=int(bh * 1.6))
        os.remove(svg)
        alle.extend(waarschuw)
    print(f"{len(prs.slides._sldIdLst)} dia's gerenderd naar {UIT}")
    if alle:
        print("\nAandachtspunten:")
        for w in alle:
            print("  -", w)
    else:
        print("Geen overloop of buitenval gevonden.")


if __name__ == "__main__":
    main()
