#!/usr/bin/env python3
"""Zet 'Acuut 2026 - procesbeschrijving strategisch capaciteitsjaarplan' om
naar het PULSE-ontwerp van 'Acute poort en Hotfloor'.

De inhoud komt een-op-een uit de aangeleverde presentatie; alleen de vorm is
opnieuw opgebouwd, met dezelfde bouwstenen als het hoofddeck zodat de dia's
er tussen passen. Alles is native PowerPoint: geen enkele afbeelding.
"""
import os

from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from deck_lib import (AMBER, BG2, BLAUW, BREED, CYAAN, DIM, GRIJS, INHOUD_Y,
                      INKT, KOL, KORAAL, MARGE, ONDER, PAPIER, RAND, RAND_OP,
                      VIOLET, _alpha, _vorm, accentpaneel, canvas, diakop,
                      haarlijn, kicker, melding, meng, paneel, penning, pijl,
                      pil, presentatie, pulslijn, rgb, spot, stroom, tabel,
                      tekst, verbinding, vlak, voet)

HIER = os.path.dirname(os.path.abspath(__file__))
UIT = os.path.join(HIER, "Procesbeschrijving-capaciteitsjaarplan.pptx")

N = {"i": 0}
DEEL = "Procesbeschrijving  ·  Strategisch capaciteitsjaarplan"


def nieuw(prs, label, titel, sub=None, kleur=CYAAN):
    d = canvas(prs)
    N["i"] += 1
    diakop(d, label, titel, sub, kleur)
    voet(d, N["i"], "Acuut 2026 - strategisch capaciteitsjaarplan")
    return d


# ------------------------------------------------------------- bouwstenen
def fasen(dia, y, items, actief, h=0.56):
    """De ketting Beddencapaciteit - Personele formatie - Financiële begroting."""
    n = len(items)
    overlap = 0.16
    b = (KOL + (n - 1) * overlap) / n
    for i, (label, kleur) in enumerate(items):
        x = MARGE + i * (b - overlap)
        v = _vorm(dia, MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON,
                  x, y, b, h)
        aan = (i == actief)
        v.fill.solid()
        v.fill.fore_color.rgb = rgb(kleur if aan else PAPIER)
        if aan:
            _alpha(v.fill.fore_color, 14)
        v.line.color.rgb = rgb(meng(kleur, BG2, 0.45) if aan else RAND)
        v.line.width = Pt(1.5 if aan else 1.0)
        v.shadow.inherit = False
        v.text_frame.text = ""
        tekst(dia, x + 0.34, y + 0.02, b - 0.62, h - 0.04,
              [{"tekst": label, "size": 12 if aan else 11.5,
                "vet": True, "kleur": kleur if aan else DIM, "na": 0,
                "uit": "center", "lh": 1.14}], anchor="midden")
    return y + h


def stapkaart(dia, x, y, w, h, nummer, titel, sub, kleur, opgetild=False):
    """Genummerde processtap als kaart."""
    if opgetild:
        accentpaneel(dia, x, y, w, h, kleur, 0.11)
    else:
        paneel(dia, x, y, w, h)
    penning(dia, x + 0.42, y + 0.36, 0.44, nummer, kleur, size=12)
    tekst(dia, x + 0.2, y + 0.64, w - 0.4, 0.48,
          [{"tekst": titel, "size": 12.5, "vet": True, "kleur": INKT, "na": 0,
            "lh": 1.16}])
    if sub:
        tekst(dia, x + 0.2, y + 1.14, w - 0.4, h - 1.26,
              [{"tekst": sub, "size": 9.5, "kleur": GRIJS, "na": 0, "lh": 1.2}])


def stapregel(dia, y, nummer, titel, sub, kleur, uitkomst=None, h=0.62):
    """Genummerde processtap als brede regel, met optionele uitkomst rechts."""
    paneel(dia, MARGE, y, KOL, h)
    penning(dia, MARGE + 0.4, y + h / 2, 0.42, nummer, kleur, size=11.5)
    uw = 3.1 if uitkomst else 0.0
    tw = 3.5
    tekst(dia, MARGE + 0.74, y + 0.04, tw, h - 0.08,
          [{"tekst": titel, "size": 12.5, "vet": True, "kleur": INKT, "na": 0,
            "lh": 1.14}], anchor="midden")
    tekst(dia, MARGE + 0.74 + tw + 0.24, y + 0.04,
          KOL - 0.98 - tw - 0.48 - uw, h - 0.08,
          [{"tekst": sub, "size": 10.5, "kleur": GRIJS, "na": 0, "lh": 1.18}],
          anchor="midden")
    if uitkomst:
        pil(dia, MARGE + KOL - 0.24 - uw, y + (h - 0.34) / 2, uw, 0.34,
            uitkomst, AMBER, size=9.5)
    return y + h


# ============================================== 1. proces bepaling capaciteit
def s_capaciteit(prs):
    d = nieuw(prs, DEEL, "Proces bepaling capaciteit",
              "Bedden en ruimtes voor het komende jaar - acht stappen van "
              "data-analyse naar een definitief advies", CYAAN)
    fasen(d, INHOUD_Y, [("Beddencapaciteit", CYAAN),
                        ("Personele formatie", DIM),
                        ("Financiële begroting", DIM)], 0)

    stappen = [
        ("Data-analyse", "instroom, patiëntaanwezigheid, ligduur en meer"),
        ("Rekenmodel opstellen", "voor komend jaar, met daarin de resultaten "
         "van de analyse op basis van de zorgvraag"),
        ("Bespreken met zorgmanager", "de resultaten uit de analyse en het "
         "rekenmodel"),
        ("Uitgangspunten ophalen", "uitgangspunten en input vanuit de "
         "zorgmanager"),
        ("Uitgangspunten verwerken", "input van de zorgmanager in het "
         "rekenmodel"),
        ("Conceptplan bespreken", "met de teammanagers"),
        ("Input teammanagers verwerken", "in het conceptplan"),
        ("Definitief advies", "met daarin het aantal benodigde bedden per "
         "week"),
    ]
    b = (KOL - 3 * 0.2) / 4
    kh = 1.55
    ry = INHOUD_Y + 0.82
    for i, (titel, sub) in enumerate(stappen):
        rij, kol = divmod(i, 4)
        x = MARGE + kol * (b + 0.2)
        y = ry + rij * (kh + 0.24)
        stapkaart(d, x, y, b, kh, f"{i + 1:02d}", titel, sub,
                  CYAAN if i < 7 else BLAUW, opgetild=(i == 7))
        if kol < 3:
            pijl(d, x + b + 0.015, y + 0.28, 0.17, 0.17, CYAAN, alpha=55)
    verbinding(d, [(MARGE + 3 * (b + 0.2) + b / 2, ry + kh),
                   (MARGE + 3 * (b + 0.2) + b / 2, ry + kh + 0.12),
                   (MARGE + b / 2, ry + kh + 0.12),
                   (MARGE + b / 2, ry + kh + 0.24)], meng(CYAAN, BG2, 0.4),
               1.25)

    melding(d, ry + 2 * kh + 0.36, "Uitkomst",
            "Het definitieve advies bevat het aantal benodigde bedden per week — "
            "het uitgangspunt voor de personele formatie.", BLAUW, 0.62)
    return d


# ================================================ 2. proces personele formatie
def s_formatie(prs):
    d = nieuw(prs, DEEL, "Proces bepaling personele formatie",
              "Van het vastgestelde aantal bedden naar een fte-berekening per "
              "afdeling", BLAUW)
    fasen(d, INHOUD_Y, [("Beddencapaciteit", DIM),
                        ("Personele formatie", BLAUW),
                        ("Financiële begroting", DIM)], 1)

    stappen = [
        ("Vastgesteld uitgangspunt", "het aantal benodigde bedden per afdeling, "
         "per week en per dag", None),
        ("Verpleegkundige normering", "dag, laat en nacht", None),
        ("Roostersleutels opstellen", "vertaling van de norm naar diensten",
         None),
        ("FTE directe uren", "berekening per week per afdeling",
         "input voor urenoverzichten"),
        ("Indirecte uren", "inventariseren en vaststellen", None),
        ("FTE-berekening per afdeling", "op basis van directe en indirecte "
         "uren", "input voor begroting P&C"),
        ("Urenoverzichten opstellen", "per afdeling, ten behoeve van de "
         "planning door CPP", None),
    ]
    y = INHOUD_Y + 0.80
    for i, (titel, sub, uit) in enumerate(stappen):
        y = stapregel(d, y, f"{i + 1:02d}", titel, sub, BLAUW, uit, 0.5) + 0.08
    return d


# ================================================= 3. proces personele planning
def s_planning(prs):
    d = nieuw(prs, DEEL, "Proces personele planning",
              "De input is bekend, het proces zelf is nog niet uitgeschreven",
              VIOLET)
    fasen(d, INHOUD_Y, [("Capaciteit\nsessies, modaliteiten en meer", VIOLET),
                        ("Personele formatie", DIM),
                        ("Financiële begroting", DIM)], 0, 0.66)

    iy = INHOUD_Y + 0.80
    kicker(d, MARGE, iy, 4.0, "wat er als input ligt", VIOLET, 10)
    invoer = [("Vastgestelde FTE-berekening",
               "inclusief roostersleutel, indirecte uren en meer"),
              ("Urenoverzicht", "per afdeling"),
              ("Opbouw Ortec", "de vertaling naar het planningssysteem")]
    b = (KOL - 2 * 0.26) / 3
    for i, (titel, sub) in enumerate(invoer):
        x = MARGE + i * (b + 0.26)
        accentpaneel(d, x, iy + 0.28, b, 1.00, VIOLET, 0.12)
        tekst(d, x + 0.26, iy + 0.38, b - 0.52, 0.80,
              [{"tekst": titel, "size": 13, "vet": True, "kleur": VIOLET,
                "na": 3},
               {"tekst": sub, "size": 10, "kleur": GRIJS, "na": 0, "lh": 1.2}],
              anchor="midden")

    py = iy + 1.46
    kicker(d, MARGE, py, 6.0, "het proces personele planning", DIM, 10)
    paneel(d, MARGE, py + 0.28, KOL, 1.24, PAPIER, RAND)
    tekst(d, MARGE + 0.3, py + 0.4, 3.2, 0.32,
          [{"tekst": "Personele planning", "size": 15, "vet": True,
            "kleur": INKT, "na": 0}], autofit=False)
    lw = (KOL - 0.6 - 5 * 0.16) / 6
    for i in range(6):
        x = MARGE + 0.3 + i * (lw + 0.16)
        v = vlak(d, x, py + 0.8, lw, 0.5, PAPIER, RAND_OP, 1.0, 0.12)
        v.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        tekst(d, x, py + 0.84, lw, 0.42,
              [{"tekst": "?", "size": 15, "vet": True, "kleur": RAND_OP,
                "na": 0, "uit": "center"}], anchor="midden", autofit=False)
        if i < 5:
            pijl(d, x + lw + 0.005, py + 0.97, 0.15, 0.15, RAND_OP, alpha=70)

    melding(d, py + 1.68, "Wat hier nog ontbreekt",
            "In de bronpresentatie staan onder 'personele planning' zes stappen "
            "nog leeg. De input is bekend, maar hoe die input via Ortec tot een "
            "rooster komt, is nog niet beschreven.", AMBER, 0.88)
    return d


# ============================================== 4. formatieberekening radiologie
def s_formatieberekening(prs):
    d = nieuw(prs, DEEL, "Formatieberekening",
              "Voorbeeld radiologie - van bruto arbeidsduur naar een fte-match",
              CYAAN)
    trappen = [("Bruto", "input vanuit HR en F&C", CYAAN),
               ("Afwezigheid", "input vanuit HR en F&C", AMBER),
               ("Netto", "bruto min afwezigheid", BLAUW),
               ("Benodigd", "input vanuit BIC", VIOLET),
               ("Resultaat", "input voor de financiële begroting", KORAAL)]
    b = (KOL - 4 * 0.22) / 5
    ty = INHOUD_Y
    for i, (naam, bron, kleur) in enumerate(trappen):
        x = MARGE + i * (b + 0.22)
        accentpaneel(d, x, ty, b, 1.10, kleur, 0.12)
        tekst(d, x + 0.2, ty + 0.12, b - 0.4, 0.36,
              [{"tekst": naam, "size": 17, "vet": True, "kleur": kleur,
                "na": 0, "uit": "center"}], autofit=False)
        haarlijn(d, x + 0.42, ty + 0.54, b - 0.84, meng(kleur, BG2, 0.5), 1)
        tekst(d, x + 0.2, ty + 0.62, b - 0.4, 0.42,
              [{"tekst": bron, "size": 9.5, "kleur": GRIJS, "na": 0,
                "uit": "center", "lh": 1.18}])
        if i < 4:
            pijl(d, x + b + 0.025, ty + 0.46, 0.17, 0.17, kleur, alpha=55)

    ly = ty + 1.36
    lw = 6.15
    paneel(d, MARGE, ly, lw, 3.18)
    tekst(d, MARGE + 0.3, ly + 0.2, lw - 0.6, 0.3,
          [{"tekst": "Rekenvoorbeeld per fte, per jaar", "size": 13.5,
            "vet": True, "kleur": INKT, "na": 0}], autofit=False)
    tekst(d, MARGE + 0.3, ly + 0.52, lw - 0.6, 0.26,
          [{"tekst": "voltijd 36,0 uur per week  ·  52,17 weken", "size": 10,
            "kleur": GRIJS, "na": 0}], autofit=False)
    rijen = [("Bruto arbeidsduur conform cao", "1.878", "100%"),
             ("Af: totale afwezigheid", "300", "15,9%"),
             ("Netto inzetbaarheid", "1.579", "84,1%")]
    tabel(d, MARGE + 0.28, ly + 0.82, lw - 0.56, [3.4, 1.2, 0.99],
          ["Onderdeel", "Uren", "Aandeel"], rijen, size=11,
          kopsize=8.5, rijhoogte=0.34, kophoogte=0.32,
          kleuren=[[INKT, INKT, GRIJS], [AMBER, AMBER, AMBER],
                   [BLAUW, BLAUW, BLAUW]],
          uitlijningen=["left", "right", "right"], kopkleur=CYAAN)

    haarlijn(d, MARGE + 0.28, ly + 2.22, lw - 0.56, RAND, 1)
    tekst(d, MARGE + 0.3, ly + 2.30, lw - 0.6, 0.24,
          [{"tekst": "WAARUIT DE AFWEZIGHEID IS OPGEBOUWD", "size": 8.5,
            "vet": True, "kleur": DIM, "na": 0, "spatie": 1.4}], autofit=False)
    posten = [("144,0", "vakantie"), ("75,1", "ziek"), ("8,0", "bijzonder"),
              ("22,0", "PLB"), ("50,4", "feestdagen")]
    pb = (lw - 0.56 - 4 * 0.1) / 5
    for i, (waarde, label) in enumerate(posten):
        px = MARGE + 0.28 + i * (pb + 0.1)
        vlak(d, px, ly + 2.56, pb, 0.5, AMBER, meng(AMBER, BG2, 0.55), 1.0,
             0.14, alpha=8)
        tekst(d, px, ly + 2.59, pb, 0.24,
              [{"tekst": waarde, "size": 12, "vet": True, "kleur": AMBER,
                "na": 0, "uit": "center"}], autofit=False)
        tekst(d, px, ly + 2.83, pb, 0.2,
              [{"tekst": label, "size": 8.5, "kleur": GRIJS, "na": 0,
                "uit": "center"}], autofit=False)

    rx = MARGE + lw + 0.34
    rw = BREED - MARGE - rx
    paneel(d, rx, ly, rw, 1.46)
    tekst(d, rx + 0.3, ly + 0.18, rw - 0.6, 0.3,
          [{"tekst": "Benodigde uren bestaan uit", "size": 13, "vet": True,
            "kleur": VIOLET, "na": 0}], autofit=False)
    onderdelen = ["indirecte uren per afdeling",
                  "directe uren op basis van roostersleutel of sessierooster",
                  "spoedmodaliteiten", "bereikbaarheidsdiensten"]
    oy = ly + 0.54
    for punt in onderdelen:
        spot(d, rx + 0.36, oy + 0.1, VIOLET, 0.1, False)
        tekst(d, rx + 0.56, oy - 0.02, rw - 0.86, 0.24,
              [{"tekst": punt, "size": 10, "kleur": GRIJS, "na": 0,
                "lh": 1.15}])
        oy += 0.22

    ay = ly + 1.60
    paneel(d, rx, ay, rw, 1.58)
    tekst(d, rx + 0.3, ay + 0.18, rw - 0.6, 0.3,
          [{"tekst": "Overige afwijkingen", "size": 13, "vet": True,
            "kleur": AMBER, "na": 0}], autofit=False)
    oy = ay + 0.50
    for punt in ["zwangerschapsverlof", "ouderschapsverlof",
                 "ziekteverzuim boven de 4 procent"]:
        spot(d, rx + 0.36, oy + 0.1, AMBER, 0.1, False)
        tekst(d, rx + 0.56, oy - 0.02, rw - 0.86, 0.24,
              [{"tekst": punt, "size": 10, "kleur": GRIJS, "na": 0,
                "lh": 1.15}])
        oy += 0.22
    tekst(d, rx + 0.3, ay + 1.16, rw - 0.6, 0.24,
          [{"tekst": "komen bovenop de begrote afwezigheid", "size": 9,
            "kleur": DIM, "na": 0}], autofit=False)

    ry2 = ly + 3.24
    accentpaneel(d, MARGE, ry2, KOL, 0.40, KORAAL, 0.10)
    vlak(d, MARGE, ry2, 0.055, 0.40, KORAAL, None, 0, 0.5)
    tekst(d, MARGE + 0.34, ry2 + 0.03, KOL - 0.68, 0.34,
          [{"tekst": "Personeelsformatie: de benodigde capaciteit in uren "
                     "afgezet tegen de bruto beschikbare formatie — het "
                     "verschil is de fte-(mis)match.", "size": 10.5,
            "kleur": INKT, "na": 0}], anchor="midden")
    return d


# ================================== 5. van patiëntaanwezigheid naar capaciteit
def s_keten(prs):
    d = nieuw(prs, DEEL, "Van patiëntaanwezigheid naar capaciteit",
              "De keten van beddencapaciteit en verpleegkundige norm naar het "
              "urenoverzicht", BLAUW)
    bronnen = [
        ("Benodigde beddencapaciteit per periode van het jaar",
         "op basis van data-analyse: patiënteninstroom, patiëntaanwezigheid "
         "en seizoenspatronen", CYAAN),
        ("Verpleegkundige norm", "per specialisme en per afdeling", VIOLET)]
    bw = (KOL - 0.36) / 2
    for i, (titel, sub, kleur) in enumerate(bronnen):
        x = MARGE + i * (bw + 0.36)
        accentpaneel(d, x, INHOUD_Y, bw, 0.92, kleur, 0.12)
        tekst(d, x + 0.28, INHOUD_Y + 0.1, bw - 0.56, 0.72,
              [{"tekst": titel, "size": 12.5, "vet": True, "kleur": kleur,
                "na": 2, "lh": 1.14},
               {"tekst": sub, "size": 9.5, "kleur": GRIJS, "na": 0,
                "lh": 1.16}], anchor="midden")

    knik = INHOUD_Y + 1.12
    mx = BREED / 2
    for i in range(2):
        x = MARGE + i * (bw + 0.36) + bw / 2
        stroom(d, x, INHOUD_Y + 0.94, mx, knik, bronnen[i][2], 3.0, alpha=55)
    spot(d, mx, knik, BLAUW, 0.14)

    stappen = [
        ("Benodigde diensten per 24 uur", "dag  ·  laat  ·  nacht  ·  overig "
         "(OVD)", BLAUW, None),
        ("Roostersleutel", "vertaling naar diensten", BLAUW, None),
        ("FTE directe zorg", "de basis voor CPP", CYAAN, "basis voor CPP"),
        ("Kortdurend ziekteverzuim", "opslag op de directe zorg", AMBER, None),
        ("Roostersleutel", "opnieuw toegepast", BLAUW, None),
        ("Overige uren", "langdurig verzuim, zwangerschap, SOMZ, PLB, "
         "geboorteverlof, overig ouderschapsverlof en overige "
         "vervolgopleidingen", AMBER, None),
        ("Dagelijkse bedrijfsvoering", "wat er in de praktijk bijkomt", VIOLET,
         None),
        ("Urenoverzicht", "de uitkomst van de keten", CYAAN, "uitkomst"),
    ]
    kb = (KOL - 3 * 0.2) / 4
    ky = knik + 0.30
    kh = 1.24
    verbinding(d, [(mx, knik), (mx, knik + 0.15),
                   (MARGE + kb / 2, knik + 0.15),
                   (MARGE + kb / 2, ky)], meng(BLAUW, BG2, 0.4), 1.25)
    for i, (titel, sub, kleur, merk) in enumerate(stappen):
        rij, kol = divmod(i, 4)
        x = MARGE + kol * (kb + 0.2)
        y = ky + rij * (kh + 0.34)
        if merk:
            accentpaneel(d, x, y, kb, kh, kleur, 0.12)
        else:
            paneel(d, x, y, kb, kh)
        vlak(d, x, y, 0.05, kh, kleur, None, 0, 0.5)
        tekst(d, x + 0.26, y + 0.16, kb - 0.5, 0.44,
              [{"tekst": titel, "size": 12.5, "vet": True, "kleur": INKT,
                "na": 0, "lh": 1.14}])
        tekst(d, x + 0.26, y + 0.62, kb - 0.5, kh - 0.78,
              [{"tekst": sub, "size": 9.5, "kleur": GRIJS, "na": 0,
                "lh": 1.18}])
        if merk:
            pil(d, x + kb - 0.26 - 1.16, y + kh - 0.34, 1.16, 0.26, merk,
                kleur, size=8.5)
        if kol < 3:
            pijl(d, x + kb + 0.015, y + kh / 2 - 0.085, 0.17, 0.17, BLAUW,
                 alpha=50)
    verbinding(d, [(MARGE + 3 * (kb + 0.2) + kb / 2, ky + kh),
                   (MARGE + 3 * (kb + 0.2) + kb / 2, ky + kh + 0.17),
                   (MARGE + kb / 2, ky + kh + 0.17),
                   (MARGE + kb / 2, ky + kh + 0.34)], meng(BLAUW, BG2, 0.4),
               1.25)

    melding(d, ky + 2 * kh + 0.44, "Hoe deze keten werkt",
            "Elke stap volgt uit de vorige: verandert de beddencapaciteit of de "
            "verpleegkundige norm, dan verandert het urenoverzicht mee.", BLAUW,
            0.56)
    return d


# ==================================================================== bouw
def bouw(pad=UIT):
    prs = presentatie()
    N["i"] = 0

    d = canvas(prs)
    N["i"] += 1
    kicker(d, MARGE, 1.62, 9.0, "Acuut 2026", CYAAN, 11)
    tekst(d, MARGE, 1.94, 9.6, 2.4,
          [{"tekst": "Procesbeschrijving", "size": 54, "vet": True,
            "kleur": INKT, "na": 0, "lh": 1.0},
           {"tekst": "strategisch capaciteitsjaarplan", "size": 54,
            "vet": True, "kleur": CYAAN, "na": 0, "lh": 1.0}], autofit=False)
    haarlijn(d, MARGE, 4.12, 3.4, CYAAN, 2.0)
    tekst(d, MARGE, 4.32, 8.6, 0.6,
          [{"tekst": "Van beddencapaciteit via personele formatie naar de "
                     "financiële begroting", "size": 15, "kleur": GRIJS,
            "na": 0}], autofit=False)
    fasen(d, 5.2, [("Beddencapaciteit", CYAAN), ("Personele formatie", BLAUW),
                   ("Financiële begroting", VIOLET)], 0, 0.6)
    tekst(d, MARGE, 5.94, KOL, 0.26,
          [{"tekst": "vier processen en een rekenvoorbeeld", "size": 10.5,
            "kleur": DIM, "na": 0, "spatie": 1.2}], autofit=False)
    pulslijn(d, 0, 6.92, BREED, 0.42, CYAAN, 1.5, 4, alpha=40)

    s_capaciteit(prs)
    s_formatie(prs)
    s_planning(prs)
    s_formatieberekening(prs)
    s_keten(prs)

    prs.save(pad)
    print("Opgeslagen:", pad, "-", len(prs.slides._sldIdLst), "dia's")
    return pad


if __name__ == "__main__":
    bouw()
