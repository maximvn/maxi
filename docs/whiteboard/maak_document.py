#!/usr/bin/env python3
"""Bouwt het Word-document. De figuren dragen de inhoud; de tekst is kort."""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HIER = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HIER, "figuren")
UIT = os.path.join(HIER, "Acute-poort-en-Hotfloor.docx")

INKT = RGBColor(0x0F, 0x1B, 0x2D)
NAVY = RGBColor(0x12, 0x3A, 0x63)
BLAUW = RGBColor(0x1F, 0x6F, 0xB2)
TEAL = RGBColor(0x0E, 0x9A, 0x87)
AMBER = RGBColor(0xDE, 0x8A, 0x1B)
ROOD = RGBColor(0xC4, 0x39, 0x2D)
PAARS = RGBColor(0x6B, 0x4B, 0xA8)
LEI = RGBColor(0x64, 0x74, 0x8B)
WIT = RGBColor(0xFF, 0xFF, 0xFF)

BREEDTE = Cm(17.2)

_NA_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
            "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
            "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
            "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
            "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
            "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def arceer(cel, hex_kleur):
    tc_pr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_kleur)
    tc_pr.append(shd)


def celmarge(tabel, boven=60, onder=60, links=120, rechts=120):
    """Celmarges in twintigsten van een punt."""
    mar = OxmlElement("w:tblCellMar")
    for kant, waarde in (("top", boven), ("left", links), ("bottom", onder),
                         ("right", rechts)):
        el = OxmlElement(f"w:{kant}")
        el.set(qn("w:w"), str(waarde))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tabel._tbl.tblPr.insert_element_before(mar, "w:tblLook", "w:tblCaption",
                                           "w:tblDescription", "w:tblPrChange")


def randen(tabel, kleur="D8E0EA", maat=4, binnen=True):
    tbl_pr = tabel._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    kanten = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if binnen else [])
    for kant in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{kant}")
        el.set(qn("w:val"), "single" if kant in kanten else "none")
        el.set(qn("w:sz"), str(maat))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), kleur)
        b.append(el)
    tbl_pr.insert_element_before(b, "w:shd", "w:tblLayout", "w:tblCellMar",
                                 "w:tblLook", "w:tblCaption", "w:tblDescription",
                                 "w:tblPrChange")


def rand(par, positie="bottom", kleur="123A63", maat=8):
    p_pr = par._p.get_or_add_pPr()
    b = p_pr.find(qn("w:pBdr"))
    if b is None:
        b = OxmlElement("w:pBdr")
        p_pr.insert_element_before(b, *_NA_PBDR)
    el = OxmlElement(f"w:{positie}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(maat))
    el.set(qn("w:space"), "4")
    el.set(qn("w:color"), kleur)
    b.append(el)


def alinea(doc, tekst="", grootte=11, kleur=INKT, vet=False, cursief=False,
           uitlijning=None, voor=0, na=8, regelafstand=1.3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(voor)
    p.paragraph_format.space_after = Pt(na)
    p.paragraph_format.line_spacing = regelafstand
    if uitlijning is not None:
        p.alignment = uitlijning
    if tekst:
        r = p.add_run(tekst)
        r.font.size = Pt(grootte)
        r.font.color.rgb = kleur
        r.bold = vet
        r.italic = cursief
    return p


def kop(doc, tekst, niveau=1, kleur=NAVY):
    h = doc.add_heading(level=niveau)
    h.paragraph_format.space_before = Pt(4 if niveau == 1 else 14)
    h.paragraph_format.space_after = Pt(2 if niveau == 1 else 4)
    r = h.add_run(tekst)
    r.font.color.rgb = kleur
    r.font.size = Pt(19 if niveau == 1 else 12.5)
    r.font.name = "Calibri"
    return h


def figuur(doc, bestand, na=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(na)
    p.add_run().add_picture(os.path.join(FIG, bestand), width=BREEDTE)
    return p


def kader(doc, label, tekst, kleur_hex="DE8A1B", vulling="FDF3E4", titelkleur=AMBER):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cel = t.cell(0, 0)
    cel.width = BREEDTE
    arceer(cel, vulling)
    p = cel.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = titelkleur
    p2 = cel.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(tekst)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = INKT
    randen(t, kleur_hex, 6, binnen=False)
    celmarge(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def vragenblok(doc, titel, kleur_hex, titelkleur, items):
    """Compacte tabel: themabalk plus genummerde vragen."""
    t = doc.add_table(rows=1 + len(items), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    kop_cel = t.cell(0, 0).merge(t.cell(0, 1))
    arceer(kop_cel, kleur_hex)
    p = kop_cel.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(titel.upper())
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = WIT
    for i, tekst in enumerate(items, start=1):
        nr, vr = t.cell(i, 0), t.cell(i, 1)
        nr.width = Cm(1.0)
        vr.width = Cm(16.2)
        if i % 2 == 0:
            arceer(nr, "F4F7FB")
            arceer(vr, "F4F7FB")
        pn = nr.paragraphs[0]
        pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pn.paragraph_format.space_before = Pt(2)
        pn.paragraph_format.space_after = Pt(2)
        rn = pn.add_run(str(i))
        rn.bold = True
        rn.font.size = Pt(9)
        rn.font.color.rgb = titelkleur
        pv = vr.paragraphs[0]
        pv.paragraph_format.space_before = Pt(2)
        pv.paragraph_format.space_after = Pt(2)
        pv.paragraph_format.line_spacing = 1.15
        rv = pv.add_run(tekst)
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = INKT
    randen(t, "D8E0EA", 4)
    celmarge(t, 40, 40, 100, 100)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def voettekst(doc):
    for sectie in doc.sections:
        p = sectie.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Acute poort en Hotfloor  -  werkdocument")
        r.font.size = Pt(8)
        r.font.color.rgb = LEI


def inhoudsopgave(doc):
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-1" \\h \\z \\u')
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Klik hier en druk op F9 om de inhoudsopgave bij te werken."
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def _herstel_zoom(doc):
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")


def hoofdstuk(doc, nummer, titel, lead, bestand, kleur=NAVY, eerste=False):
    if not eerste:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    kop(doc, f"{nummer}.  {titel}", 1, kleur)
    p = alinea(doc, lead, 11, LEI, na=4)
    rand(p, "bottom", "D8E0EA", 4)
    figuur(doc, bestand)


def bouw():
    doc = Document()
    _herstel_zoom(doc)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.font.color.rgb = INKT

    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.9)
    s.top_margin = Cm(1.8)
    s.bottom_margin = Cm(1.6)
    voettekst(doc)

    # ============================================================= omslag
    for _ in range(3):
        doc.add_paragraph()
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    cel = band.cell(0, 0)
    cel.width = BREEDTE
    arceer(cel, "123A63")
    p = cel.paragraphs[0]
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("INTEGRAAL CAPACITEITSMANAGEMENT")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x9E, 0xC0, 0xE0)
    p2 = cel.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    r2 = p2.add_run("Acute poort en Hotfloor")
    r2.bold = True
    r2.font.size = Pt(34)
    r2.font.color.rgb = WIT
    p3 = cel.add_paragraph()
    p3.paragraph_format.space_after = Pt(28)
    r3 = p3.add_run("Van de huidige situatie naar de nieuwbouw, juni 2027")
    r3.font.size = Pt(13)
    r3.font.color.rgb = RGBColor(0xC9, 0xDC, 0xEE)
    randen(band, "123A63", 4, binnen=False)
    celmarge(band, 0, 0, 400, 400)

    doc.add_paragraph()
    alinea(doc, "Dit document vat de whiteboardsessie samen in twaalf platen. Elke "
                "plaat is op zichzelf te lezen: wat weten we, wat staat er nog open, "
                "en wat moet er besloten worden.", 11.5, LEI, na=16)

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate([
            ("Onderwerp", "Acute poort, Hotfloor, scenario's en personele inzet"),
            ("Afdelingen", "ICU (10 kamers), CCU/SCU/EHH (15 kamers), SEH, kind spoed"),
            ("Bron", "Whiteboardsessie, aangevuld met de vragen van Lonneke en Maxim"),
            ("Status", "Werkdocument - de gemarkeerde punten staan nog open")]):
        meta.cell(i, 0).width = Cm(3.4)
        meta.cell(i, 1).width = Cm(13.8)
        arceer(meta.cell(i, 0), "F4F7FB")
        rk = meta.cell(i, 0).paragraphs[0].add_run(k)
        rk.bold = True
        rk.font.size = Pt(9.5)
        rk.font.color.rgb = NAVY
        rv = meta.cell(i, 1).paragraphs[0].add_run(v)
        rv.font.size = Pt(9.5)
    randen(meta, "D8E0EA", 4)
    celmarge(meta, 60, 60, 120, 120)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ======================================================== inhoud
    kop(doc, "Inhoud", 1, NAVY)
    inhoudsopgave(doc)
    alinea(doc, "Klik in de inhoudsopgave en druk op F9 om de paginanummers te vullen.",
           9, LEI, cursief=True, na=18)

    kop(doc, "Afkortingen", 2, NAVY)
    afk = doc.add_table(rows=5, cols=2)
    afk.alignment = WD_TABLE_ALIGNMENT.CENTER
    paren = [("SEH", "spoedeisende hulp"), ("EHH", "eerste hart hulp"),
             ("ICU", "intensive care unit"), ("CCU", "coronary care unit"),
             ("SCU", "stroke care unit"), ("vpk", "verpleegkundige"),
             ("CPP", "capaciteits- en personeelsplanning"),
             ("BIC", "business intelligence centrum"),
             ("D / T / L / N", "dag-, tussen-, laat- en nachtdienst"),
             ("EPA", "entrustable professional activity")]
    for i in range(5):
        for j in range(2):
            k, v = paren[i + j * 5]
            cel = afk.cell(i, j)
            cel.width = Cm(8.6)
            p = cel.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            rk = p.add_run(k + "   ")
            rk.bold = True
            rk.font.size = Pt(9.5)
            rk.font.color.rgb = NAVY
            rv = p.add_run(v)
            rv.font.size = Pt(9.5)
            rv.font.color.rgb = LEI
    randen(afk, "D8E0EA", 4)
    celmarge(afk, 40, 40, 120, 120)

    # ============================================================ 1 t/m 10
    hoofdstuk(doc, 1, "Startsituatie per afdeling",
              "Voordat er gerekend kan worden moet per afdeling vastliggen wat de norm "
              "is. Op dit moment ligt die voor geen enkele afdeling vast.",
              "fig_01_startsituatie.png", NAVY)
    kader(doc, "Let op bij de aantallen",
          "10 en 15 zijn de kamers die ICU en CCU/SCU/EHH nu hebben. Het zijn "
          "verschillende kamers voor verschillende specialismen, dus ze zijn niet bij "
          "elkaar op te tellen. De berekeningen zijn per specialisme apart gemaakt, op "
          "de huidige gescheiden situatie; in de nieuwbouwscenario's waarin afdelingen "
          "samengaan moet opnieuw geteld worden. Bij de SEH is niet af te lezen of er "
          "een jaarplan ligt, en bij kind spoed ontbreken de gegevens vrijwel helemaal.",
          "C4392D", "FBEDEB", ROOD)

    hoofdstuk(doc, 2, "Tijdpad",
              "Het rekenwerk start in september 2026. Van oktober 2026 tot en met mei "
              "2027 werken de afdelingen in de oudbouw; in juni 2027 wordt er verhuisd.",
              "fig_02_tijdpad.png", PAARS)

    hoofdstuk(doc, 3, "Berekeningen",
              "De berekening bestaat uit drie onderdelen. Het aantal bedden wordt "
              "bepaald op basis van patientaanwezigheid, per afdeling uitgesplitst naar "
              "stroom.", "fig_03_berekeningen.png", PAARS)

    hoofdstuk(doc, 4, "Van vier locaties naar twee clusters",
              "De acute en intensieve zorg ligt nu op vier plaatsen in het gebouw en "
              "gaat naar twee clusters: de acute poort en de Hotfloor.",
              "fig_04_oudnieuw.png", BLAUW)

    hoofdstuk(doc, 5, "De nieuwbouw fysiek",
              "Van de nieuwbouw staat alleen het aantal bedden op de Hotfloor vast. "
              "Voor de SEH en de EHH is het aantal plekken nog niet bepaald.",
              "fig_05_nieuwbouw.png", TEAL)

    hoofdstuk(doc, 6, "Vijf scenario's",
              "Er liggen vijf mogelijke varianten. Elke variant verandert zowel de "
              "fysieke capaciteit als de personele inzet.",
              "fig_06_scenarios.png", BLAUW)

    hoofdstuk(doc, 7, "Van scenario naar rooster",
              "Zodra de keuze gemaakt is, ligt de route naar het rooster vast: van "
              "week- en dagplan via de roostersleutels naar de planning en de monitoring.",
              "fig_07_stappenplan.png", PAARS)

    hoofdstuk(doc, 8, "Analyse 1: de acute poort",
              "De analyse is gedaan en gedeeld. De uitkomst is dat een acute poort voor "
              "kind, EHH en SEH samen niet past.",
              "fig_08_acutepoort.png", BLAUW)

    hoofdstuk(doc, 9, "Analyse 2: de Hotfloor",
              "De Hotfloor past wel op zestien bedden. Twee punten staan nog open, en "
              "de EHH-vraag raakt beide wijzigingen.",
              "fig_09_hotfloor.png", TEAL)

    hoofdstuk(doc, 10, "Wat weten we per situatie?",
              "Per situatie is in beeld gebracht wat er bekend is over de fysieke "
              "planning en over de personele inzet.",
              "fig_10_matrix.png", NAVY)

    # ================================================================ 11
    hoofdstuk(doc, 11, "De vragen van Lonneke en Maxim",
              "Zesendertig vragen en knelpunten, gegroepeerd naar het moment waarop ze "
              "beantwoord moeten worden.",
              "fig_11_vragenoverzicht.png", AMBER)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    kop(doc, "11.1  Uitgangspunten: oudbouw en transitie", 2, BLAUW)
    vragenblok(doc, "Oudbouw en transitie", "1F6FB2", BLAUW, [
        "Wat moeten we meenemen uit lopende projecten, of wordt dit later doorgevoerd "
        "in de begroting en het jaarplan?",
        "Hoe werken de afdelingen nu al toe naar de nieuwbouwsituatie, en welke impact "
        "heeft dat op bedden en personeel?",
        "Kunnen we de huidige personele inzet van de SEH aanhouden tot de nieuwbouw, "
        "exclusief eventuele extra inzet voor de nieuwbouw?",
        "De bedbezetting is in het weekend lager dan doordeweeks; formatie daarop "
        "aanpassen komt neer op 5-5-4. In de oudbouw geen aanpassing, alleen toewerken "
        "naar de nieuwbouw: motiveren en dubbel opleiden.",
        "Voor de ICU houden we de landelijke norm aan, maar daar zit bij ons ook "
        "recovery in. Welke normering hanteren we? (Geparkeerd voor later.)",
        "Wat is het uitgangspunt voor de AO-dienst? Bij drukte inzet in de zorg; negen "
        "van de tien keer blijf je buiten de zorg.",
        "Per september start de regieverpleegkundige. Wat betekent dat voor de "
        "personele inzet en voor de indirecte uren, en staat die functie binnen of "
        "buiten de zorg?",
    ])

    kop(doc, "11.2  Uitgangspunten: verhuisperiode", 2, AMBER)
    vragenblok(doc, "Verhuisperiode", "DE8A1B", AMBER, [
        "Welke extra diensten zijn er nodig tijdens de verhuisperiode ten opzichte van "
        "de oudbouw, en op welke locatie worden die ingezet?",
    ])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    kop(doc, "11.3  Uitgangspunten: nieuwbouw", 2, TEAL)
    vragenblok(doc, "Nieuwbouw", "0E9A87", TEAL, [
        "Waar komt de personele inzet op de EHH vandaan: van de SEH of van de ICU/CCU?",
        "Dat wordt bepaald door het afkapmoment van de acute cardiologiepatient van SEH "
        "naar EHH/CCU. In de berekeningen is uitgegaan van inzet vanaf de CCU en een "
        "afkapmoment op 2 uur; nu wordt gekeken naar scenario's op 4, 6 en 24 uur "
        "waarbij de EHH-bedden bij de SEH zijn meegenomen.",
        "Doet Remco een uitspraak over de verpleegkundige norm, of houden we de huidige "
        "normen aan en leggen we die aan hem voor?",
        "In de nieuwbouw wordt de scope (CCU) anders ingericht. Hoe verwerken we dat in "
        "de personele inzet? Nu is er een extra dienst overdag en extra belasting in de "
        "late en de nachtdienst.",
        "De ICU heeft een normenkader waarbij maximaal 10% afgeweken mag worden van "
        "dedicated inzet van ICU-verpleegkundigen.",
        "Is er al nagedacht over de personele inzet op de SEH? Kunnen we het patroon "
        "opplussen op basis van de extra bedden, en van hoeveel bedden gaan we uit?",
        "Indirecte uren: de oude begroting loopt sterk uiteen. Wat houden we aan, en "
        "verlaagt de regieverpleegkundige dit per september?",
        "Hoe richten we de inzet van de regieverpleegkundige in de nieuwbouw in?",
        "Komt de kinderverpleegkundige dedicated op de SEH, of op afroep?",
        "Hoe staat het met de EPA's? Zijn er vereisten voor SEH-verpleegkundigen, en "
        "zijn er verplichte EPA's voor ICU en CCU?",
        "Als die EPA's er zijn, wat betekent dat dan in de praktijk?",
        "Zetten we voor de nieuwbouw nog in op een acute poule voor CCU/ICU en CCU/SEH? "
        "Nu is dat alleen ICU/SEH.",
    ])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    kop(doc, "11.4  Knelpunten", 2, ROOD)
    vragenblok(doc, "Knelpunten", "C4392D", ROOD, [
        "Verpleegkundige norm: aanbodgericht werken in plaats van vraaggestuurd.",
        "SEH-artsen: de inzet is verhoogd maar de begroting is niet aangepast.",
        "Acute poort: wat komt waar, en wat kunnen wij doorrekenen?",
        "ICU en CCU wijken structureel af van het jaarplan BIC.",
        "De personele planning loopt niet volledig via CPP, met name rond de acute poule.",
        "Indirecte uren.",
        "Uitgangspunten en kaders ontbreken.",
        "Vakantiegoedkeuringen niet integraal en te ruim vrijgegeven, met name bij de "
        "CCU. Actie: samen met Remco een vakantiegoedkeuring opstellen; hij noemde 20 "
        "tot 25%.",
        "Slapers op de SEH.",
    ])
    kader(doc, "Kanttekening bij de JDT",
          "Ervaren werkdruk: je kunt vol liggen waardoor de druk hoog is, terwijl de "
          "JDT-score dat niet laat zien.", "64748B", "F4F7FB", LEI)

    kop(doc, "11.5  Simulatie en jaarplan", 2, PAARS)
    vragenblok(doc, "Simulatie nieuwbouw", "6B4BA8", PAARS, [
        "Volledige ICU-data.",
        "CCU-data exclusief de eerste 2 uur van CCU-patienten die naar de SEH gaan.",
        "Op hoeveel bedden komen we op dagniveau uit, en zijn er seizoenspatronen? Hoe "
        "is de verdeling ICU/CCU voor het personeel?",
        "Versneld jaarplan CCU, en het document van Daniek: huidig, transitie, "
        "nieuwbouw.",
    ])

    kop(doc, "11.6  Aantekeningen aftrap jaarplannen Acuut 2027", 2, NAVY)
    vragenblok(doc, "Aantekeningen 24-7-2026", "123A63", NAVY, [
        "We kijken nu naar de verpleegkundige inzet; voor de SEH-artsen is eerder al "
        "iets gedaan. Wat is de wens en de verwachting?",
        "Inzet op integratie van afdelingen per nu: ICU en CCU, waarbij de SEH nog even "
        "buiten beschouwing blijft.",
        "Er is nog geen inzet van de ICU op de CCU; andersom gebeurt dat wel, maar "
        "alleen ad hoc.",
    ])

    kop(doc, "11.7  Direct naar Remco", 2, AMBER)
    vragenblok(doc, "Op de mail naar Remco", "DE8A1B", AMBER, [
        "Hoort de EHH bij het personeel van de ICU/CCU of bij de SEH?",
        "Waar is de opvang van de acute cardiologiepatient: op de SEH of op de EHH?",
        "Hoe lang blijft een patient op de SEH voordat die naar de EHH gaat?",
    ])
    kader(doc, "Masterplan Acuut",
          "Wat ontbreekt er nog voor de transitie naar de nieuwbouw? Wat hebben we "
          "nodig op basis van patientaanwezigheid? Voorstel: een jaarplan opstellen en "
          "dat naast het huidige plan leggen. Alle punten, met alle disciplines erbij, "
          "op de mail zetten.", "123A63", "F4F7FB", NAVY)

    # ================================================================ 12
    hoofdstuk(doc, 12, "Wat moet er besloten worden",
              "Zes beslissingen waar de rest van het traject op wacht. De eerste drie "
              "blokkeren de onderbouwing van de personele inzet.",
              "fig_12_besluiten.png", ROOD)

    doc.save(UIT)
    print("Document opgeslagen:", UIT)
    return UIT


if __name__ == "__main__":
    bouw()
