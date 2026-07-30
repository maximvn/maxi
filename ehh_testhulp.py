# -*- coding: utf-8 -*-
"""Gedeelde hulpfuncties voor de tests van het EHH-model."""
import subprocess
import sys

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

SP = "/tmp/claude-0/-home-user-maxi/6cb3d4b7-2b09-53b1-b52a-d878fb334a3e/scratchpad"
RECALC = "/root/.claude/skills/xlsx/scripts/recalc.py"
BUILD = "/home/user/maxi/build_ehh_analyse.py"

# berekende kolommen op het tabblad Data
C_INSEL, C_STATUS, C_LIG, C_LIGTXT = "AD", "AE", "AF", "AG"
C_START, C_EIND, C_BEST = "AH", "AI", "AJ"
C_W4, C_W6, C_W24 = "AK", "AL", "AM"
C_EERSTE, C_LIGSEL, C_WINSTSEL, C_CTRL = "AN", "AO", "AP", "AQ"
CALC_KOLOMMEN = [C_INSEL, C_STATUS, C_LIG, C_LIGTXT, C_START, C_EIND, C_BEST,
                 C_W4, C_W6, C_W24, C_EERSTE, C_LIGSEL, C_WINSTSEL, C_CTRL,
                 "AR", "AS"]

# instelcellen
I_BRON, I_ONDER, I_BOVEN = "C6", "C11", "C12"
I_DEDUP = "C15"
I_BESTKOL, I_BESTWAARDE, I_TELBEST = "C19", "C21", "C22"
I_FILTER_AAN, I_FILTERKOL, I_FILTERZOEK, I_FILTERWAARDE = "C25", "C26", "C28", "C29"
I_SCENARIO, I_TIJDJA, I_EXCL = "C33", "C34", "C35"

# dashboardcellen
D_REGELS, D_METFORMULE, D_UNIEK, D_GELDIG, D_MELDING = "C6", "C7", "C8", "C9", "B10"
D_SEL, D_HUIS, D_BLIJFT, D_EERDER = "C13", "C14", "C15", "C16"
D_WINST_UREN, D_WINST_GEM, D_WINST_MED = "E17", "E18", "E19"
D_KORTSTE, D_MEDIAAN, D_GEMIDDELD, D_LANGSTE = "C22", "C23", "C24", "C25"
D_SCEN4, D_SCEN6, D_SCEN24 = 29, 30, 31
D_STATUS_EERSTE, D_STATUS_TOTAAL = 34, 42

STATUSSEN = ["In selectie - naar huis", "In selectie - blijft in ziekenhuis",
             "Te kort (t/m ondergrens)", "Te lang (boven bovengrens)",
             "Geen geldige start- of eindtijd", "Extra regel van dezelfde opname",
             "Buiten het extra filter", "Tijd onbetrouwbaar (uitgesloten)"]


def bouw(pad, rijen=60):
    subprocess.run([sys.executable, BUILD], check=True,
                   env={"EHH_ROWS": str(rijen), "EHH_OUT": pad,
                        "PATH": "/usr/bin:/bin"})


def trek_formules_door(pad, tot_rij):
    """Bootst na wat de gebruiker doet: rij 2 kopieren naar beneden."""
    wb = load_workbook(pad)
    ws = wb["Data"]
    bron = {k: ws[f"{k}2"].value for k in CALC_KOLOMMEN}
    opmaak = {k: ws[f"{k}2"].number_format for k in CALC_KOLOMMEN}
    for r in range(3, tot_rij + 1):
        for k in CALC_KOLOMMEN:
            cel = ws[f"{k}{r}"]
            cel.value = Translator(bron[k], origin=f"{k}2").translate_formula(f"{k}{r}")
            cel.number_format = opmaak[k]
    wb.save(pad)


def herbereken(pad, fouten, timeout="240"):
    r = subprocess.run([sys.executable, RECALC, pad, timeout],
                       capture_output=True, text=True)
    if '"total_errors": 0' not in r.stdout:
        fouten.append(f"formulefouten in {pad}: {r.stdout}{r.stderr}")
    return r.stdout


def num(ws, cel):
    v = ws[cel].value
    return 0 if v in (None, "") else v


def maak_check(fouten):
    def check(naam, got, verw, tol=0.05):
        try:
            afwijking = abs(float(got) - float(verw))
        except (TypeError, ValueError):
            fouten.append(f"{naam}: kreeg {got!r}, verwacht {verw!r}")
            return
        if afwijking > tol:
            fouten.append(f"{naam}: {got} verwacht {verw}")
    return check
