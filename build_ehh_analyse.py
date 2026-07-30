# -*- coding: utf-8 -*-
"""
Bouwt 'EHH ligduur analyse.xlsx' - bewust simpel gehouden.

Kern
----
Ligduur = OpnameperiodeDatumTijdTotEnMet (kolom H) minus
          OpnameperiodeDatumTijdVan     (kolom G), maal 24 = uren.

Er zijn drie hulpkolommen met DIRECTE celverwijzingen (=$H2-$G2 e.d.), zodat
het doortrekken van rij 2 naar 250.000 rijen gewoon werkt. Geen instellingen-
tabblad, geen INDEX over de hele rij, geen named ranges, geen shared formulas,
geen voorwaardelijke opmaak die naar een ander tabblad verwijst - dat waren de
oorzaken van eerdere problemen.
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment

N = int(os.environ.get("EHH_ROWS", "250000"))
LAST = N + 1
UIT = os.environ.get("EHH_OUT", "/home/user/maxi/EHH ligduur analyse.xlsx")
VOORGEVULD = int(os.environ.get("EHH_PREFILL", "1"))   # aantal rijen met formule
FONT = "Arial"

BRON_KOPPEN = [
    "OpnameIDIntern", "CodePreferent", "Jaar",
    "OpnameDatumTijdVan", "OpnameDatumTijdTotEnMet",
    "OpnameperiodeIDIntern",
    "OpnameperiodeDatumTijdVan", "OpnameperiodeDatumTijdTotEnMet",
    "AfdelingCode", "KamerCode", "Code", "IsSpoedOpname",
    "SpecialismeCodeIntern", "OpnameDiagnoseOmschrijving", "OmschrijvingIntern",
    "OP_AantalOpnamePeriode", "OP_DoorlooptijdOpnamePeriodeInMinuten",
    "Cardioversie?", "Overname?", "Rijnstate",
    "Datum tijd van", "Datum tot en met",
    "Ontslagdatumtijd Scenario 1 -  4 uur",
    "Ontslagdatumtijd Scenario 2 -  6 uur",
    "Ontslagdatumtijd Scenario 3 -  24 uur",
    "Patientvolgorde", "CCU-Eerst?", "Tijd bekend",
]
assert len(BRON_KOPPEN) == 28

# kolommen G en H bevatten start en eind (OpnameperiodeDatumTijd van / tot)
KOL_START = "G"
KOL_EIND = "H"
ONDER, BOVEN = 6, 24

# hulpkolommen direct na de brondata (AC, AD, AE)
C_LIG = "AC"
C_FLAG = "AD"
C_SEL = "AE"
HELPERS = [
    (C_LIG, "Ligduur (uren)", 14, "getal2"),
    (C_FLAG, f"Tussen {ONDER} en {BOVEN} uur", 15, None),
    (C_SEL, f"Ligduur als {ONDER}-{BOVEN} uur", 16, "getal2"),
]

K_KOP, K_HELP = "1F3864", "375623"
f_titel = Font(name=FONT, size=16, bold=True, color=K_KOP)
f_sub = Font(name=FONT, size=11, italic=True, color="595959")
f_kop = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_blokkop = Font(name=FONT, size=12, bold=True, color=K_KOP)
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_kpi = Font(name=FONT, size=16, bold=True, color=K_KOP)
f_waarsch = Font(name=FONT, size=11, bold=True, color="C00000")

fill_kop = PatternFill("solid", fgColor=K_KOP)
fill_help = PatternFill("solid", fgColor=K_HELP)
fill_blok = PatternFill("solid", fgColor="D9E2F3")
fill_groen = PatternFill("solid", fgColor="E2EFDA")
fill_invul = PatternFill("solid", fgColor="FFF2CC")
dun = Side(style="thin", color="BFBFBF")
rand = Border(left=dun, right=dun, top=dun, bottom=dun)
FMT = {"datum": "dd-mm-yyyy hh:mm", "getal0": "#,##0", "getal2": "#,##0.00",
       "proc": "0.0%"}


def zet(ws, cel, waarde, font=f_norm, fill=None, fmt=None, align=None,
        wrap=False, border=False):
    c = ws[cel]
    c.value = waarde
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = FMT.get(fmt, fmt)
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if border:
        c.border = rand
    return c


wb = Workbook()
wb.remove(wb.active)

# ====================================================================== DATA ==
ws_d = wb.create_sheet("Data")
# alleen de kopregel en kolom A vastzetten (niet meer op AD: dan bevriezen alle
# kolommen ervoor en kun je niet meer naar rechts scrollen)
ws_d.freeze_panes = "B2"

for i, kop in enumerate(BRON_KOPPEN, start=1):
    c = ws_d.cell(row=1, column=i, value=kop)
    c.font, c.fill = f_kop, fill_kop
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_d.column_dimensions[get_column_letter(i)].width = max(12, min(len(kop)+2, 30))
for letter, titel, br, _ in HELPERS:
    c = ws_d[f"{letter}1"]
    c.value, c.font, c.fill = titel, f_kop, fill_help
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_d.column_dimensions[letter].width = br
ws_d.row_dimensions[1].height = 44

ws_d["A1"].comment = Comment(
    "Plak uw data vanaf cel A2 (zonder de kopregel), via Plakken > Waarden.\n"
    "Plak niet over de gekleurde kolommen AC/AD/AE heen.", "Model")
ws_d[f"{C_LIG}1"].comment = Comment(
    "Ligduur = kolom H (OpnameperiodeDatumTijdTotEnMet) minus kolom G "
    "(OpnameperiodeDatumTijdVan), in uren.\n\n"
    "De formules staan alleen in rij 2. Trek ze zelf door:\n"
    f"1. selecteer {C_LIG}2:{C_SEL}2\n"
    "2. Ctrl+C\n"
    f"3. selecteer {C_LIG}3:{C_SEL}<uw laatste rij>  (Ctrl+Shift+Pijl-omlaag)\n"
    "4. Ctrl+V", "Model")


def formules(r):
    lig = f"${C_LIG}{r}"
    flag = f"${C_FLAG}{r}"
    return {
        # ligduur in uren; leeg als een van beide tijden ontbreekt. De IFERROR
        # vangt het geval af waarin G/H tekst blijken in plaats van een datum:
        # dan wordt het leeg in plaats van #WAARDE!, en het Dashboard wijst er
        # via 'Regels met een geldige ligduur' op.
        C_LIG: (f'=IF(OR(${KOL_START}{r}="",${KOL_EIND}{r}=""),"",'
                f'IFERROR((${KOL_EIND}{r}-${KOL_START}{r})*24,""))'),
        # Ja als de ligduur groter is dan de ondergrens en niet groter dan de
        # bovengrens; anders Nee
        C_FLAG: (f'=IF({lig}="","",'
                 f'IF(AND({lig}>{ONDER},{lig}<={BOVEN}),"Ja","Nee"))'),
        # de ligduur, maar alleen voor de regels tussen 6 en 24 uur (voor de
        # samenvatting op het Dashboard)
        C_SEL: f'=IF({flag}="Ja",{lig},"")',
    }


fmt_map = {h[0]: h[3] for h in HELPERS}
for r in range(2, 2 + max(1, VOORGEVULD)):
    for letter, expr in formules(r).items():
        c = ws_d[f"{letter}{r}"]
        c.value = expr
        c.font = f_norm
        if fmt_map.get(letter):
            c.number_format = FMT[fmt_map[letter]]

ws_d.auto_filter.ref = f"A1:{C_SEL}{LAST}"
# voorwaardelijke opmaak met LETTERLIJKE tekst (geen verwijzing naar een ander
# tabblad): kleurt de regels die tussen 6 en 24 uur vallen groen
ws_d.conditional_formatting.add(
    f"{C_FLAG}2:{C_FLAG}{LAST}",
    CellIsRule(operator="equal", formula=['"Ja"'], fill=fill_groen))

# ================================================================== DASHBOARD =
ws_db = wb.create_sheet("Dashboard")
ws_db.sheet_view.showGridLines = False
for col, br in [("A", 3), ("B", 44), ("C", 16), ("D", 13), ("E", 40)]:
    ws_db.column_dimensions[col].width = br


def R(k):
    return f"Data!${k}$2:${k}${LAST}"


AC, AD, AE = R(C_LIG), R(C_FLAG), R(C_SEL)

zet(ws_db, "B2", "Dashboard EHH ligduur", f_titel)
zet(ws_db, "B3",
    f"Kernvraag: hoeveel regels hebben een ligduur langer dan {ONDER} en tot en "
    f"met {BOVEN} uur.", f_sub)

# controle vooraf
INGELEZEN = f'COUNTA({R("A")})'
MET_FORMULE = f'COUNTA({AC})'      # telt ook formulecellen die "" teruggeven
MET_LIG = f'COUNT({AC})'           # telt alleen echte getallen (geldige ligduur)
zet(ws_db, "B5", "Controle", f_blokkop, fill_blok, border=True)
zet(ws_db, "C5", "Aantal", f_bold, fill_blok, align="center", border=True)
zet(ws_db, "E5", "Toelichting", f_bold, fill_blok, align="center", border=True)
zet(ws_db, "B6", "Regels ingelezen", f_norm, border=True, align="left")
zet(ws_db, "C6", f"={INGELEZEN}", f_bold, fmt="getal0", align="center", border=True)
zet(ws_db, "E6", "Nul = nog geen data geplakt.", f_sub, border=True, wrap=True,
    align="left")
zet(ws_db, "B7", "Regels waar de formules staan", f_norm, border=True, align="left")
zet(ws_db, "C7", f"={MET_FORMULE}", f_bold, fmt="getal0", align="center",
    border=True)
zet(ws_db, "E7",
    "Lager dan 'Regels ingelezen'? Dan zijn de formules nog niet ver genoeg "
    "doorgetrokken (zie Handleiding stap 2).",
    f_sub, border=True, wrap=True, align="left")
zet(ws_db, "B8", "Regels met een geldige ligduur", f_norm, border=True, align="left")
zet(ws_db, "C8", f"={MET_LIG}", f_bold, fmt="getal0", align="center", border=True)
zet(ws_db, "E8",
    "Kleiner dan 'Regels waar de formules staan'? Dan mist bij die regels een "
    "datum in kolom G of H, of staat daar tekst in plaats van een datum.",
    f_sub, border=True, wrap=True, align="left")
zet(ws_db, "B9",
    f'=IF({INGELEZEN}=0,"Nog geen data geplakt op het tabblad Data.",'
    f'IF({MET_FORMULE}<{INGELEZEN},'
    f'"LET OP: er zijn "&{INGELEZEN}&" regels, maar de formules staan pas op "&'
    f'{MET_FORMULE}&" regels. Trek {C_LIG}2:{C_SEL}2 verder naar beneden.",'
    f'"In orde - alle regels zijn doorgerekend."))',
    f_waarsch, fill_invul, border=True, wrap=True, align="left")
ws_db.row_dimensions[9].height = 30

# kernantwoord
zet(ws_db, "B11", f"Tussen {ONDER} en {BOVEN} uur", f_blokkop, fill_blok, border=True)
zet(ws_db, "C11", "Aantal", f_bold, fill_blok, align="center", border=True)
zet(ws_db, "D11", "% van geldige", f_bold, fill_blok, align="center", border=True)
zet(ws_db, "E11", "Toelichting", f_bold, fill_blok, align="center", border=True)
TUSSEN = f'COUNTIF({AD},"Ja")'
zet(ws_db, "B12", f"Regels met ligduur > {ONDER} en <= {BOVEN} uur", f_bold,
    fill_groen, border=True, align="left")
zet(ws_db, "C12", f"={TUSSEN}", f_kpi, fill_groen, fmt="getal0", align="center",
    border=True)
zet(ws_db, "D12", f'=IF({MET_LIG}=0,"",{TUSSEN}/{MET_LIG})', f_bold, fill_groen,
    fmt="proc", align="center", border=True)
zet(ws_db, "E12", "DIT IS HET GEVRAAGDE AANTAL.", f_sub, fill_groen, border=True,
    wrap=True, align="left")

# ligduurstatistiek binnen de selectie
zet(ws_db, "B14", f"Ligduur binnen {ONDER}-{BOVEN} uur", f_blokkop, fill_blok,
    border=True)
zet(ws_db, "C14", "Uren", f_bold, fill_blok, align="center", border=True)
r = 15
for label, fn in (("Kortste ligduur", "MIN"), ("Gemiddelde ligduur", "AVERAGE"),
                  ("Mediane ligduur", "MEDIAN"), ("Langste ligduur", "MAX")):
    zet(ws_db, f"B{r}", label, f_norm, border=True, align="left")
    zet(ws_db, f"C{r}", f'=IF(COUNT({AE})=0,"",ROUND({fn}({AE}),2))', f_norm,
        fmt="getal2", align="center", border=True)
    r += 1

# verdeling naar ligduur
r += 1
zet(ws_db, f"B{r}", "Verdeling naar ligduur (alle regels met ligduur)", f_blokkop,
    fill_blok, border=True)
zet(ws_db, f"C{r}", "Aantal", f_bold, fill_blok, align="center", border=True)
zet(ws_db, f"D{r}", "% van geldige", f_bold, fill_blok, align="center", border=True)
r += 1
blokken = [(0, 2), (2, 4), (4, 6), (6, 8), (8, 10), (10, 12), (12, 14),
           (14, 16), (16, 18), (18, 20), (20, 22), (22, 24)]
for lo, hi in blokken:
    crit = f'{AC},">"&{lo},{AC},"<="&{hi}'
    label = f"{lo} tot {hi} uur"
    fill = fill_groen if lo >= ONDER and hi <= BOVEN else None
    zet(ws_db, f"B{r}", label, f_norm, fill, border=True, align="left")
    zet(ws_db, f"C{r}", f'=COUNTIFS({crit})', f_norm, fill, fmt="getal0",
        align="center", border=True)
    zet(ws_db, f"D{r}", f'=IF({MET_LIG}=0,"",COUNTIFS({crit})/{MET_LIG})', f_norm,
        fill, fmt="proc", align="center", border=True)
    r += 1
zet(ws_db, f"B{r}", "Meer dan 24 uur", f_norm, border=True, align="left")
zet(ws_db, f"C{r}", f'=COUNTIFS({AC},">24")', f_norm, fmt="getal0",
    align="center", border=True)
zet(ws_db, f"D{r}", f'=IF({MET_LIG}=0,"",COUNTIFS({AC},">24")/{MET_LIG})', f_norm,
    fmt="proc", align="center", border=True)
r += 1
zet(ws_db, f"B{r}", "Totaal met ligduur", f_bold, fill_blok, border=True,
    align="left")
zet(ws_db, f"C{r}", f"=SUM(C{r-len(blokken)-1}:C{r-1})", f_bold, fill_blok,
    fmt="getal0", align="center", border=True)
zet(ws_db, f"D{r}", "", f_norm, fill_blok, border=True)

# =============================================================== HANDLEIDING ==
ws_h = wb.create_sheet("Handleiding")
ws_h.sheet_view.showGridLines = False
for col, br in [("A", 3), ("B", 30), ("C", 96)]:
    ws_h.column_dimensions[col].width = br
zet(ws_h, "B2", "EHH ligduur analyse", f_titel)
zet(ws_h, "B3", "Ligduur op basis van OpnameperiodeDatumTijd van en tot en met.",
    f_sub)

REGELS = [
    ("STAP 1 - Data plakken", ""),
    ("", "Ga naar tabblad 'Data'. Rij 1 bevat exact uw 28 kolomnamen in dezelfde volgorde."),
    ("", "Klik op cel A2 en plak uw dataset zonder de kopregel, via Plakken > Waarden."),
    ("", "Plak niet over de gekleurde kolommen AC, AD en AE heen."),
    ("", ""),
    ("STAP 2 - Formules doortrekken  (BELANGRIJK)", ""),
    ("", "De formules staan alleen in rij 2, zodat het bestand klein blijft en Excel"),
    ("", "het zonder foutmelding opent. Trek ze zelf door naar beneden:"),
    ("", f"1. selecteer {C_LIG}2 tot en met {C_SEL}2"),
    ("", "2. Ctrl+C (kopieren)"),
    ("", f"3. selecteer {C_LIG}3 tot en met {C_SEL} op uw laatste datarij"),
    ("", "   (tip: Ctrl+Shift+Pijl-omlaag selecteert razendsnel tot onderaan)"),
    ("", "4. Ctrl+V (plakken)"),
    ("", "Blok 'Controle' op het Dashboard meldt of u ver genoeg bent gekomen."),
    ("", ""),
    ("WAT DE KOLOMMEN DOEN", ""),
    ("Ligduur (uren)  [AC]",
     "Kolom H (OpnameperiodeDatumTijdTotEnMet) min kolom G "
     "(OpnameperiodeDatumTijdVan), in uren."),
    (f"Tussen {ONDER} en {BOVEN} uur  [AD]",
     f"'Ja' als de ligduur groter is dan {ONDER} en niet groter dan {BOVEN} uur, "
     f"anders 'Nee'."),
    (f"Ligduur als {ONDER}-{BOVEN} uur  [AE]",
     "Dezelfde ligduur, maar alleen ingevuld bij de regels met 'Ja'. "
     "Het Dashboard gebruikt deze kolom voor de kortste/gemiddelde/langste ligduur."),
    ("", ""),
    ("STAP 3 - Lijst op kortste ligduur", ""),
    ("", "Op 'Data': klik op het filterpijltje van kolom AD en vink alleen 'Ja' aan."),
    ("", "Klik daarna op het filterpijltje van 'Ligduur (uren)' (AC) en kies"),
    ("", "Sorteren van klein naar groot. U ziet dan de selectie, kortste ligduur bovenaan."),
    ("", ""),
    ("STAP 4 - Antwoord", ""),
    ("", "Het tabblad 'Dashboard' toont het aantal regels tussen 6 en 24 uur,"),
    ("", "plus de kortste, gemiddelde en langste ligduur en een verdeling per 2 uur."),
    ("", ""),
    ("ALS DE LIGDUUR LEEG OF #WAARDE! IS", ""),
    ("", "Dan worden de kolommen G en H als tekst gelezen in plaats van als datum."),
    ("", "Selecteer kolom G en H, ga naar Gegevens > Tekst naar kolommen > Voltooien;"),
    ("", "Excel maakt er dan echte datums van en de ligduur verschijnt vanzelf."),
]
r = 5
for kop, tekst in REGELS:
    if kop and not tekst:
        zet(ws_h, f"B{r}", kop, f_blokkop)
    elif kop:
        zet(ws_h, f"B{r}", kop, f_bold, wrap=True, align="left")
        zet(ws_h, f"C{r}", tekst, f_norm, wrap=True, align="left")
    else:
        zet(ws_h, f"C{r}", tekst, f_norm, wrap=True, align="left")
    r += 1

wb.move_sheet("Handleiding", offset=-(len(wb.sheetnames) - 1))
wb.active = 0
wb.save(UIT)
print(f"opgeslagen: {UIT} | bereik tot rij {LAST} | formulerijen {VOORGEVULD}")
