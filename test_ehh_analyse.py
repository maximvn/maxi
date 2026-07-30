# -*- coding: utf-8 -*-
"""Vult een kleine testset in het model en controleert de uitkomsten."""
import datetime as dt
import subprocess
import sys

from openpyxl import load_workbook

PAD = "/tmp/claude-0/-home-user-maxi/6cb3d4b7-2b09-53b1-b52a-d878fb334a3e/scratchpad/test.xlsx"

# id, afdeling, start, ligduur_uren, overname, opname_eind_extra_min, tijd_bekend
CASES = [
    # binnen selectie, naar huis -> categorie B
    ("P01", "EHH", "2025-01-02 08:00", 6.5, "", 0, "Ja"),
    ("P02", "EHH", "2025-01-02 09:00", 12.0, "", 0, "Ja"),
    ("P03", "EHH", "2025-01-03 10:00", 23.5, "", 0, "Ja"),
    ("P04", "EHH", "2025-01-03 11:00", 6.01, "", 0, "Ja"),
    # binnen selectie, blijft in ziekenhuis -> categorie C
    ("P05", "EHH", "2025-01-04 07:00", 8.0, "Ja", 0, "Ja"),
    ("P06", "EHH", "2025-01-04 12:00", 20.0, "", 600, "Ja"),
    # onder de ondergrens -> A
    ("P07", "EHH", "2025-01-05 08:00", 3.0, "", 0, "Ja"),
    ("P08", "EHH", "2025-01-05 09:00", 6.0, "", 0, "Ja"),   # exact 6u -> A
    # boven de bovengrens -> D
    ("P09", "EHH", "2025-01-06 08:00", 30.0, "", 0, "Ja"),
    # niet EHH -> E
    ("P10", "CCU", "2025-01-06 08:00", 10.0, "", 0, "Ja"),
    # geen datum -> F
    ("P11", "EHH", "", 0, "", 0, "Ja"),
]


def bouw():
    subprocess.run([sys.executable, "/home/user/maxi/build_ehh_analyse.py"],
                   check=True, env={"EHH_ROWS": "40", "EHH_OUT": PAD,
                                    "PATH": "/usr/bin:/bin"})


def vul():
    wb = load_workbook(PAD)
    ws = wb["Data"]
    for i, (pid, afd, start, uren, ovn, extra, tb) in enumerate(CASES):
        r = 2 + i
        ws[f"A{r}"] = pid
        ws[f"C{r}"] = 2025
        ws[f"I{r}"] = afd
        ws[f"J{r}"] = f"K{i+1:02d}"
        ws[f"L{r}"] = "Ja"
        ws[f"M{r}"] = "CAR"
        ws[f"N{r}"] = "Pijn op de borst"
        ws[f"S{r}"] = ovn
        ws[f"AB{r}"] = tb
        if not start:
            continue
        s = dt.datetime.strptime(start, "%Y-%m-%d %H:%M")
        e = s + dt.timedelta(hours=uren)
        ws[f"G{r}"] = s                       # OpnameperiodeDatumTijdVan
        ws[f"H{r}"] = e                       # OpnameperiodeDatumTijdTotEnMet
        ws[f"D{r}"] = s                       # OpnameDatumTijdVan
        ws[f"E{r}"] = e + dt.timedelta(minutes=extra)
        ws[f"U{r}"] = s
        ws[f"V{r}"] = e
        ws[f"Q{r}"] = round(uren * 60)
        # scenariokolommen: afgekapte ontslagtijd
        for kol, cap in (("W", 4), ("X", 6), ("Y", 24)):
            ws[f"{kol}{r}"] = min(e, s + dt.timedelta(hours=cap))
        for kol in ("G", "H", "D", "E", "U", "V", "W", "X", "Y"):
            ws[f"{kol}{r}"].number_format = "dd-mm-yyyy hh:mm"
    wb.save(PAD)


def controleer():
    wb = load_workbook(PAD, data_only=True)
    d, db = wb["Data"], wb["Dashboard"]

    verwacht_cat = {
        "P01": "B", "P02": "B", "P03": "B", "P04": "B",
        "P05": "C", "P06": "C",
        "P07": "A", "P08": "A",
        "P09": "D", "P10": "E", "P11": "F",
    }
    fouten = []
    for i, (pid, *_) in enumerate(CASES):
        r = 2 + i
        cat = (d[f"AO{r}"].value or "")
        if not cat.startswith(verwacht_cat[pid]):
            fouten.append(f"{pid}: categorie '{cat}' verwacht {verwacht_cat[pid]}")

    def num(ws, cel):
        v = ws[cel].value
        return 0 if v in (None, "") else v

    # verwachte kerncijfers
    sel = [c for c in CASES if verwacht_cat[c[0]] in ("B", "C")]
    huis = [c for c in CASES if verwacht_cat[c[0]] == "B"]
    winst6 = sum(round(c[3] - 6, 2) for c in huis)

    checks = [
        ("selectie 6-24u", "C10", len(sel)),
        ("naar huis", "C11", len(huis)),
        ("blijft in ziekenhuis", "C12", len(sel) - len(huis)),
        ("had eerder gekund", "C13", len(huis)),
        ("totale winst uren", "E14", round(winst6, 1)),
    ]
    for naam, cel, verw in checks:
        got = num(db, cel)
        if abs(float(got) - float(verw)) > 0.05:
            fouten.append(f"Dashboard {cel} ({naam}): {got} verwacht {verw}")

    # sortering: oplopend op ligduur
    ws_s = wb["Selectie"]
    ligduren = []
    for r in range(6, 6 + len(sel)):
        v = ws_s[f"L{r}"].value
        if v not in (None, ""):
            ligduren.append(float(v))
    if len(ligduren) != len(sel):
        fouten.append(f"Selectie: {len(ligduren)} rijen gevuld, verwacht {len(sel)}")
    if ligduren != sorted(ligduren):
        fouten.append(f"Selectie niet oplopend gesorteerd: {ligduren}")

    # controlekolom ligduur vs doorlooptijd moet 0 zijn
    for i, (pid, _, start, *_) in enumerate(CASES):
        if not start:
            continue
        v = d[f"BA{2+i}"].value
        if v not in (None, "") and abs(v) > 0.5:
            fouten.append(f"{pid}: controle doorlooptijd wijkt af ({v} min)")

    print("ligduren in Selectie:", ligduren)
    print("Dashboard C10..C13:", [num(db, f"C{r}") for r in (10, 11, 12, 13)],
          "winst E14:", num(db, "E14"))
    if fouten:
        print("\nFOUTEN:")
        for f in fouten:
            print(" -", f)
        sys.exit(1)
    print("\nAlle controles geslaagd.")


if __name__ == "__main__":
    bouw()
    vul()
    res = subprocess.run(
        [sys.executable, "/root/.claude/skills/xlsx/scripts/recalc.py", PAD, "180"],
        capture_output=True, text=True)
    print(res.stdout.strip())
    controleer()
